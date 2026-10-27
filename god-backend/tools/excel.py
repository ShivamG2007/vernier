"""Preview/import the team's workbook, or export the combined database.

Historical outbox rows are imported as sent so a data import cannot send old alerts.
"""

import argparse
import json
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
DEFAULT_SOURCE = os.environ.get("VERNIER_DATABASE_XLSX", str(ROOT / "Vernier_Database.xlsx"))

import store  # noqa: E402

HEADERS = {
    "lost_found": ["id", "kind", "item", "colour", "place", "details", "contact", "source", "status", "created"],
    "events": ["id", "title", "category", "date", "venue", "details", "major", "source", "created"],
    "listings": ["id", "kind", "item", "price", "details", "contact", "status", "created"],
    "machines": ["id", "location", "status", "updated"],
    "bookings": ["id", "room", "day", "start", "end", "booked_by"],
    "students": ["contact", "interests", "level"],
    "outbox": ["id", "to_chat", "text", "sent", "created"],
}


def workbook_rows(path):
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("Excel tools need openpyxl. Install god-backend/requirements.txt.") from None
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for name, columns in HEADERS.items():
        if name not in wb.sheetnames:
            raise ValueError(f"Missing sheet: {name}")
        rows = wb[name].iter_rows(values_only=True)
        header = [str(v).strip() if v is not None else "" for v in next(rows)]
        if header[:len(columns)] != columns:
            raise ValueError(f"Unexpected columns in {name}: {header[:len(columns)]}")
        for number, values in enumerate(rows, start=2):
            values = values[:len(columns)]
            if any(v is not None and str(v).strip() for v in values):
                yield name, number, dict(zip(columns, values))


def as_text(value, limit=4000):
    return str(value).strip()[:limit] if value is not None and str(value).strip() else None


def insert_legacy(con, table, row):
    if table == "lost_found":
        kind = as_text(row["kind"])
        item = as_text(row["item"], 120)
        if kind not in {"lost", "found"} or not item:
            return None
        source = as_text(row["source"]) if row["source"] in {"chat", "camera"} else "chat"
        raw_status = as_text(row["status"])
        status = ("claimed" if raw_status == "claimed" else "unclaimed") if kind == "found" else (
            "matched" if raw_status == "matched" else "open")
        return con.execute("""INSERT INTO items
            (kind,item,color,place,custom_string,source,reporter,status,alerts_enabled,created_at)
            VALUES (?,?,?,?,?,?,?,?,0,?)""",
            (kind, item, as_text(row["colour"], 60), as_text(row["place"], 120),
             as_text(row["details"]), source, as_text(row["contact"], 200) or "",
             status, as_text(row["created"]) or store.now())).lastrowid
    if table == "events":
        if not as_text(row["title"]):
            return None
        return con.execute("""INSERT INTO events
            (title,category,date,venue,details,major,source,created_at)
            VALUES (?,?,?,?,?,?,?,?)""",
            (as_text(row["title"], 160), as_text(row["category"], 40) or "other",
             as_text(row["date"], 100), as_text(row["venue"], 120), as_text(row["details"]),
             int(bool(row["major"])), as_text(row["source"], 200),
             as_text(row["created"]) or store.now())).lastrowid
    if table == "listings":
        kind = as_text(row["kind"])
        if kind not in {"buy", "sell"} or not as_text(row["item"]):
            return None
        return con.execute("""INSERT INTO listings
            (kind,item,price,details,contact,status,alerts_enabled,created_at) VALUES (?,?,?,?,?,?,0,?)""",
            (kind, as_text(row["item"], 120), as_text(row["price"], 80), as_text(row["details"]),
             as_text(row["contact"], 200), as_text(row["status"]) or "open",
             as_text(row["created"]) or store.now())).lastrowid
    if table == "machines":
        machine_id = as_text(row["id"], 80)
        if not machine_id:
            return None
        con.execute("""INSERT OR IGNORE INTO machines(id,location,status,updated_at)
            VALUES (?,?,?,?)""", (machine_id, as_text(row["location"], 160) or "Unknown",
                                 as_text(row["status"], 80) or "unknown",
                                 as_text(row["updated"]) or store.now()))
        return machine_id
    if table == "bookings":
        if not all(as_text(row[k]) for k in ("room", "day", "start", "end")):
            return None
        return con.execute("""INSERT INTO bookings
            (room,day,start,end,booked_by,created_at) VALUES (?,?,?,?,?,?)""",
            (as_text(row["room"], 120), as_text(row["day"], 10), as_text(row["start"], 5),
             as_text(row["end"], 5), as_text(row["booked_by"], 200) or "",
             store.now())).lastrowid
    if table == "students":
        contact = as_text(row["contact"], 200)
        if not contact:
            return None
        interests = [s.strip().casefold() for s in (as_text(row["interests"]) or "").split(",") if s.strip()]
        con.execute("INSERT OR IGNORE INTO students(contact,interests,level,active) VALUES (?,?,?,0)",
                    (contact, json.dumps(interests), "major" if row["level"] == "major" else "all"))
        return contact
    if table == "outbox":
        if not as_text(row["to_chat"]) or not as_text(row["text"]):
            return None
        return con.execute("""INSERT INTO outbox
            (to_chat,text,state,created_at,sent_at) VALUES (?,?,'sent',?,?)""",
            (as_text(row["to_chat"], 200), as_text(row["text"], 2000),
             as_text(row["created"]) or store.now(), store.now())).lastrowid
    raise ValueError(f"Unknown table: {table}")


def import_workbook(source, apply):
    entries = list(workbook_rows(source))
    counts = {name: 0 for name in HEADERS}
    for name, _, _ in entries:
        counts[name] += 1
    print("Workbook rows:", counts)
    if not apply:
        print("Preview only. Add --apply to import into the combined database.")
        return
    store.init_db()
    imported = 0
    with store.connect() as con:
        for name, row_number, row in entries:
            source_key = str(row.get("id") or row.get("contact") or row_number)
            exists = con.execute("""SELECT 1 FROM legacy_imports WHERE
                source_name=? AND table_name=? AND source_key=?""",
                (Path(source).name, name, source_key)).fetchone()
            if exists:
                continue
            local_id = insert_legacy(con, name, row)
            if local_id is None:
                continue
            con.execute("""INSERT INTO legacy_imports(source_name,table_name,source_key,local_id)
                VALUES (?,?,?,?)""", (Path(source).name, name, source_key, str(local_id)))
            imported += 1
    print(f"Imported {imported} new rows into {store.DB_PATH}. Historical outbox rows were marked sent.")


def export_workbook(destination):
    try:
        import openpyxl
    except ImportError:
        raise SystemExit("Excel tools need openpyxl. Install god-backend/requirements.txt.") from None
    store.init_db()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    queries = {
        "lost_found": "SELECT id,kind,item,color AS colour,place,custom_string AS details,reporter AS contact,source,status,created_at AS created FROM items",
        "events": "SELECT id,title,category,date,venue,details,major,source,created_at AS created FROM events",
        "listings": "SELECT id,kind,item,price,details,contact,status,created_at AS created FROM listings",
        "machines": "SELECT id,location,status,updated_at AS updated FROM machines",
        "bookings": "SELECT id,room,day,start,end,booked_by FROM bookings",
        "students": "SELECT contact,interests,level FROM students",
        "outbox": "SELECT id,to_chat,text,CASE WHEN state='sent' THEN 1 ELSE 0 END AS sent,created_at AS created FROM outbox",
    }
    with store.connect() as con:
        for name, headers in HEADERS.items():
            sheet = wb.create_sheet(name)
            sheet.append(headers)
            for row in con.execute(queries[name]):
                values = list(row)
                if name == "students":
                    values[1] = ",".join(json.loads(values[1]))
                sheet.append(values)
    destination = Path(destination).resolve()
    if destination == Path(DEFAULT_SOURCE).resolve():
        raise ValueError("Export cannot overwrite the team's source workbook")
    wb.save(destination)
    print(f"Exported to {destination}")


def main():
    parser = argparse.ArgumentParser(description="Team workbook import/export")
    sub = parser.add_subparsers(dest="command", required=True)
    imp = sub.add_parser("import")
    imp.add_argument("--source", default=DEFAULT_SOURCE)
    imp.add_argument("--apply", action="store_true")
    exp = sub.add_parser("export")
    exp.add_argument("--destination", default=str(ROOT / "vernier_export.xlsx"))
    check = sub.add_parser("check")
    check.add_argument("--source", default=DEFAULT_SOURCE)
    args = parser.parse_args()
    if args.command == "import":
        import_workbook(args.source, args.apply)
    elif args.command == "export":
        export_workbook(args.destination)
    else:
        counts = {name: 0 for name in HEADERS}
        for name, _, _ in workbook_rows(args.source):
            counts[name] += 1
        print("Workbook columns match the expected team schema. Rows:", counts)


if __name__ == "__main__":
    main()
