import os
import sys

# Ensure backend root is in sys.path
backend_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

try:
    import db
except ImportError as e:
    print(f"Error importing db: {e}")
    sys.exit(1)

try:
    import openpyxl
except ImportError:
    print("Error: openpyxl is not installed. Please install it (e.g. pip install openpyxl).")
    sys.exit(1)

TARGET_TABLES = ["lost_found", "events", "listings", "machines", "bookings", "students", "outbox"]


def run_import():
    force = "--force" in sys.argv
    db_target = os.path.abspath(db.DB_PATH)
    db_filename = os.path.basename(db_target).lower()

    if db_filename == "vernier.db" and not force:
        print(
            f"ERROR: Refusing to run against '{db.DB_PATH}' (vernier.db).\n"
            "To confirm writing to this database, rerun with --force:\n"
            "    python tools/import_xlsx.py --force"
        )
        sys.exit(1)

    excel_path = os.getenv(
        "VERNIER_DATABASE_XLSX",
        os.path.join(backend_dir, "Vernier_Database.xlsx")
    )

    if not os.path.exists(excel_path):
        print(f"Error: Excel file not found at {excel_path}")
        sys.exit(1)

    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
    except Exception as e:
        print(f"Error opening workbook {excel_path}: {e}")
        sys.exit(1)

    # Ensure tables exist
    db.init_db()

    counts = {}

    for table in TARGET_TABLES:
        # Find sheet name (case-insensitive fallback)
        sheet_name = None
        for name in wb.sheetnames:
            if name.lower() == table.lower():
                sheet_name = name
                break

        if not sheet_name:
            print(f"Warning: Sheet '{table}' not found in {excel_path}, skipping.")
            counts[table] = 0
            continue

        sheet = wb[sheet_name]

        # Header row 1
        headers = []
        for col_idx in range(1, sheet.max_column + 1):
            val = sheet.cell(row=1, column=col_idx).value
            if val is not None and str(val).strip():
                headers.append((col_idx, str(val).strip()))

        if not headers:
            counts[table] = 0
            continue

        header_names = [h[1] for h in headers]
        col_indices = [h[0] for h in headers]
        placeholders = ", ".join(["?"] * len(header_names))
        columns_clause = ", ".join(f'"{col}"' for col in header_names)
        sql = f'INSERT OR REPLACE INTO "{table}" ({columns_clause}) VALUES ({placeholders})'

        inserted_count = 0
        # Data from row 2, stop at the first fully empty row
        for row_idx in range(2, sheet.max_row + 1):
            raw_vals = [sheet.cell(row=row_idx, column=c).value for c in col_indices]

            # Check if entire row is empty
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in raw_vals):
                break

            # Convert empty string/whitespace to None (SQL NULL)
            row_data = []
            for v in raw_vals:
                if v is None:
                    row_data.append(None)
                elif isinstance(v, str) and not v.strip():
                    row_data.append(None)
                else:
                    row_data.append(v)

            db.execute(sql, tuple(row_data))
            inserted_count += 1

        counts[table] = inserted_count

    print(f"\nImport successful into: {db.DB_PATH}")
    print("-" * 35)
    for table, count in counts.items():
        print(f"  {table:<15}: {count:>4} rows")
    print("-" * 35)
    print(f"  {'Total':<15}: {sum(counts.values()):>4} rows\n")


if __name__ == "__main__":
    run_import()
