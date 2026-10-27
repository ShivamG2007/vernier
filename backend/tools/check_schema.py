import os
import sys

# Ensure backend root is in sys.path so we can import db
backend_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if backend_dir not in sys.path:
    sys.path.insert(0, backend_dir)

try:
    import db
except ImportError as e:
    print(f"Error importing db: {e}")
    sys.exit(1)

try:
    import openpyxl
except ImportError:
    print("Error: openpyxl is not installed. Please install it (e.g. pip install openpyxl).")
    sys.exit(1)


def check_schema():
    excel_path = os.getenv(
        "VERNIER_DATABASE_XLSX",
        os.path.join(backend_dir, "Vernier_Database.xlsx")
    )

    if not os.path.exists(excel_path):
        print(f"Error: Excel file not found at {excel_path}")
        sys.exit(1)

    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
    except Exception as e:
        print(f"Error loading workbook {excel_path}: {e}")
        sys.exit(1)

    if "Data_Dictionary" not in wb.sheetnames:
        print(f"Error: Sheet 'Data_Dictionary' not found. Available sheets: {wb.sheetnames}")
        sys.exit(1)

    sheet = wb["Data_Dictionary"]

    header_row = 4
    headers = {}
    for col in range(1, sheet.max_column + 1):
        val = sheet.cell(row=header_row, column=col).value
        if val is not None:
            headers[str(val).strip().lower()] = col

    def find_col(candidates):
        for c in candidates:
            for h, idx in headers.items():
                if c == h:
                    return idx
        for c in candidates:
            for h, idx in headers.items():
                if c in h:
                    return idx
        return None

    col_table = find_col(["table", "table name"])
    col_column = find_col(["column", "column name", "field"])
    col_type = find_col(["sqlite type", "data type", "type"])
    col_key = find_col(["key", "primary key", "pk"])

    if not col_table or not col_column:
        print(f"Error: Could not identify required Table/Column headers on row 4. Found: {list(headers.keys())}")
        sys.exit(1)

    expected_tables = {}
    current_table = None

    for r in range(header_row + 1, sheet.max_row + 1):
        t_val = sheet.cell(row=r, column=col_table).value if col_table else None
        c_val = sheet.cell(row=r, column=col_column).value if col_column else None
        type_val = sheet.cell(row=r, column=col_type).value if col_type else None
        key_val = sheet.cell(row=r, column=col_key).value if col_key else None

        if t_val is not None and str(t_val).strip():
            current_table = str(t_val).strip()

        if not current_table or c_val is None or not str(c_val).strip():
            continue

        col_name = str(c_val).strip()
        exp_type = str(type_val).strip() if type_val is not None else ""
        key_raw = str(key_val).strip() if key_val is not None else ""
        is_pk = "PK" in key_raw.upper() or "PRIMARY" in key_raw.upper()

        if current_table not in expected_tables:
            expected_tables[current_table] = {}

        expected_tables[current_table][col_name] = {
            "type": exp_type,
            "is_pk": is_pk,
            "key_raw": key_raw,
        }

    # Initialize the database using db.init_db()
    db.init_db()

    con = db.connect()
    try:
        db_tables = [r["name"] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()]
        db_tables_map = {t.lower(): t for t in db_tables}

        mismatches = []

        for exp_table, exp_cols in expected_tables.items():
            tbl_key = exp_table.lower()
            if tbl_key not in db_tables_map:
                mismatches.append(f"Missing table: '{exp_table}'")
                continue

            actual_table_name = db_tables_map[tbl_key]
            pragma_rows = con.execute(f'PRAGMA table_info("{actual_table_name}")').fetchall()
            actual_cols_map = {r["name"].lower(): r for r in pragma_rows}

            for exp_col, meta in exp_cols.items():
                col_key = exp_col.lower()
                if col_key not in actual_cols_map:
                    mismatches.append(f"Missing column: '{exp_table}.{exp_col}'")
                    continue

                act_col = actual_cols_map[col_key]
                act_type = (act_col["type"] or "").strip()
                act_pk = bool(act_col["pk"])

                exp_type = meta["type"]
                if exp_type and exp_type.upper() != act_type.upper():
                    mismatches.append(
                        f"Wrong type on '{exp_table}.{exp_col}': expected '{exp_type}', got '{act_type}'"
                    )

                if meta["is_pk"] != act_pk:
                    exp_pk_desc = "PRIMARY KEY" if meta["is_pk"] else "NOT PRIMARY KEY"
                    act_pk_desc = "PRIMARY KEY" if act_pk else "NOT PRIMARY KEY"
                    mismatches.append(
                        f"Wrong primary key on '{exp_table}.{exp_col}': expected {exp_pk_desc}, got {act_pk_desc}"
                    )
    finally:
        con.close()

    if mismatches:
        print(f"SCHEMA MISMATCHES ({len(mismatches)}):")
        for m in mismatches:
            print(f"  - {m}")
        sys.exit(1)
    else:
        print("OK")
        sys.exit(0)


if __name__ == "__main__":
    check_schema()
